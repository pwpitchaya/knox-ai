"""Page-preserving extraction and chunking for evidence-backed retrieval."""
from pathlib import Path

# Mappings checked against rendered Thai glyphs in this pilot corpus. Unknown
# private-use glyphs still require OCR; this is not a universal font decoder.
LEGACY_THAI = str.maketrans({'\uf70a': '่', '\uf70b': '้', '\uf70c': '๊',
                           '\uf70d': '๋', '\uf70e': '์', '\uf712': '็', '\uf703': 'ึ'})


def extract_sections(file_path: str, file_type: str) -> list[dict]:
    if file_type == "pdf":
        from pypdf import PdfReader
        sections = []
        for number, page in enumerate(PdfReader(file_path).pages, start=1):
            raw = page.extract_text() or ""
            text = raw.translate(LEGACY_THAI)
            # Missing text or legacy Thai private-use glyphs need image OCR;
            # never guess a font-specific character mapping.
            damaged = (any('\ue000' <= char <= '\uf8ff' for char in text)
                       or '\ufffd' in text)
            if len(text.strip()) < 15 or damaged:
                from ocr import available, cached_ocr
                if available():
                    sections.append(cached_ocr(file_path, number))
                    continue
            sections.append({"text": text, "page_number": number,
                             "method": "pdf-text-normalized" if text != raw else "pdf-text",
                             "needs_text_review": damaged or text != raw})
        return sections
    if file_type == "docx":
        from docx import Document
        document = Document(file_path)
        # Word pagination depends on rendering; never invent a page number.
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        paragraphs.extend(" | ".join(cell.text for cell in row.cells)
                          for table in document.tables for row in table.rows)
        return [{"text": "\n".join(paragraphs), "page_number": None}]
    if file_type == "txt":
        raw = Path(file_path).read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("tis-620")
        return [{"text": text, "page_number": None}]
    raise ValueError("Unsupported file type")


def chunk_sections(sections: list[dict], chunk_size: int = 800,
                   overlap: int = 120) -> list[dict]:
    if not 0 <= overlap < chunk_size:
        raise ValueError("Require 0 <= overlap < chunk_size")
    chunks = []
    for section in sections:
        text = section["text"].strip()
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            fragment = text[start:end].strip()
            if fragment:
                chunks.append({"text": fragment,
                               "page_number": section["page_number"]})
            if end == len(text):
                break
            start = end - overlap
    return chunks
